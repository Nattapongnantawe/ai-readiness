"""Accuracy of the LLM extraction (rubric: >= 85%), checked by a person.

Step 1 - sample jobs into an Excel sheet to check by hand:
    python -m scripts.eval_extract sample              # 50 jobs (25 per site) -> output/eval_extract.xlsx
    python -m scripts.eval_extract sample --n 100

Step 2 - open the sheet, read each job's text and fill the yellow columns:
    hard_wrong     extracted hard skills that are wrong        (copy names, separate with ;)
    hard_missing   hard skills in the text that were missed    (separate with ;)
    soft_wrong / soft_missing   same for soft skills
    salary_ok / level_ok / work_mode_ok   Y = correct, N = wrong, blank = cannot tell
  Leave a cell empty when nothing is wrong / missing.

Step 3 - score it:
    python -m scripts.eval_extract score               # -> output/eval_result.txt

Precision = correct / extracted, Recall = correct / (correct + missing), F1 = 2PR / (P + R),
micro-averaged over all sampled jobs. Contact details in the text are masked (PDPA); the sheet still
comes from the live database, so keep it on this computer.
"""
import argparse
import re
import sys
from pathlib import Path

from sqlalchemy import create_engine, text

from app.config import settings
from app.privacy import clean_title, mask_text

OUT = Path("output")
XLSX = OUT / "eval_extract.xlsx"
FILL = ["hard_wrong", "hard_missing", "soft_wrong", "soft_missing", "salary_ok", "level_ok", "work_mode_ok"]


def _openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        sys.exit("ต้องติดตั้งก่อน:  pip install openpyxl")


def split(s) -> list[str]:
    return [x.strip() for x in re.split(r"[;\n]", str(s or "")) if x.strip()]


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip().lower()


def cmd_sample(a) -> None:
    xl = _openpyxl()
    eng = create_engine(settings.database_url)
    with eng.connect() as c:
        cols = {r[0] for r in c.execute(text(
            "SELECT column_name FROM information_schema.columns WHERE table_name = 'jobs'"))}
        sal_txt = next((x for x in ("salary_text", "salary_raw", "salary") if x in cols), None)
        c.execute(text("SELECT setseed(:s)"), {"s": a.seed})
        jobs = c.execute(text(f"""
            SELECT * FROM (
              SELECT j.id, lower(j.source_site::text) AS source, j.title,
                     coalesce(j.detail_url, j.source_url) AS url,
                     j.description, j.qualification,
                     {f"j.{sal_txt}::text" if sal_txt else "NULL"} AS salary_text,
                     j.salary_min, j.salary_max, j.seniority::text AS level, j.work_mode::text AS work_mode,
                     row_number() OVER (PARTITION BY j.source_site ORDER BY random()) AS rn
              FROM jobs j WHERE j.llm_extracted_at IS NOT NULL) t
            WHERE rn <= :per ORDER BY source, id"""), {"per": a.n // 2}).mappings().all()
        ids = [j["id"] for j in jobs]
        sk = c.execute(text("""
            SELECT js.job_id, lower(s.skill_type::text) AS t, s.name
            FROM job_skills js JOIN skills s ON s.id = js.skill_id
            WHERE js.job_id = ANY(:ids) ORDER BY s.name"""), {"ids": ids}).mappings().all()

    skills: dict[int, dict[str, list[str]]] = {}
    for r in sk:
        kind = "soft" if r["t"] and r["t"].startswith("soft") else "hard"
        skills.setdefault(r["job_id"], {"hard": [], "soft": []})[kind].append(r["name"])

    wb = xl.Workbook()
    ws = wb.active
    ws.title = "check"
    head = ["job_id", "source", "title", "url", "text (description + qualification)", "salary_text",
            "salary_min", "salary_max", "level", "work_mode", "hard_extracted", "soft_extracted"] + FILL
    ws.append(head)
    for j in jobs:
        s = skills.get(j["id"], {"hard": [], "soft": []})
        body = "\n\n".join(x for x in (mask_text(j["description"]), mask_text(j["qualification"])) if x)
        ws.append([j["id"], j["source"], clean_title(j["title"]), j["url"], body[:30000], j["salary_text"],
                   j["salary_min"], j["salary_max"], j["level"], j["work_mode"],
                   "; ".join(s["hard"]), "; ".join(s["soft"])] + [None] * len(FILL))

    yellow = xl.styles.PatternFill("solid", fgColor="FFF2CC")
    wrap = xl.styles.Alignment(wrap_text=True, vertical="top")
    widths = {"A": 8, "B": 9, "C": 30, "D": 14, "E": 70, "F": 16, "K": 34, "L": 24}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for i in range(len(head) - len(FILL) + 1, len(head) + 1):
        letter = xl.utils.get_column_letter(i)
        ws.column_dimensions[letter].width = 18
        for cell in ws[letter]:
            cell.fill = yellow
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    ws.freeze_panes = "C2"
    OUT.mkdir(exist_ok=True)
    wb.save(XLSX)
    print(f"✅ {len(jobs)} งาน -> {XLSX}")
    print("กรอกช่องสีเหลือง แล้วรัน: python -m scripts.eval_extract score")


def prf(tp: int, fp: int, fn: int) -> tuple[float, float, float]:
    p = tp / (tp + fp) if tp + fp else 0.0
    r = tp / (tp + fn) if tp + fn else 0.0
    return p, r, (2 * p * r / (p + r) if p + r else 0.0)


def cmd_score(a) -> None:
    xl = _openpyxl()
    if not XLSX.exists():
        sys.exit(f"ไม่พบ {XLSX} — รัน sample ก่อน")
    wb = xl.load_workbook(XLSX)
    ws = wb["check"] if "check" in wb.sheetnames else wb.active
    who = "ตรวจโดยผู้จัดทำ"
    if "criteria" in wb.sheetnames and wb["criteria"]["A1"].value == "annotator":
        who = wb["criteria"]["B1"].value or who
    rows = list(ws.iter_rows(values_only=True))
    head, data = list(rows[0]), [dict(zip(rows[0], r)) for r in rows[1:] if r and r[0]]

    tot = {k: [0, 0, 0] for k in ("hard", "soft")}
    warn = []
    for r in data:
        for k in ("hard", "soft"):
            ext = {norm(x) for x in split(r[f"{k}_extracted"])}
            wrong = {norm(x) for x in split(r[f"{k}_wrong"])}
            miss = {norm(x) for x in split(r[f"{k}_missing"])}
            if wrong - ext:
                warn.append(f"job {r['job_id']}: {k}_wrong ไม่อยู่ในรายการที่สกัด: {', '.join(sorted(wrong - ext))}")
            wrong &= ext
            tot[k][0] += len(ext) - len(wrong)
            tot[k][1] += len(wrong)
            tot[k][2] += len(miss)

    lines = [f"ประเมินการสกัดข้อมูล: {len(data)} งาน", f"ผู้ติดป้าย: {who}", ""]
    lines.append(f"{'':14}{'Precision':>10}{'Recall':>10}{'F1':>10}{'ถูก':>7}{'ผิด':>7}{'ขาด':>7}")
    allk = [sum(tot[k][i] for k in tot) for i in range(3)]
    for label, (tp, fp, fn) in [("Hard skills", tot["hard"]), ("Soft skills", tot["soft"]), ("ทักษะรวม", allk)]:
        p, rc, f = prf(tp, fp, fn)
        lines.append(f"{label:14}{p:10.1%}{rc:10.1%}{f:10.1%}{tp:7}{fp:7}{fn:7}")
    lines.append("")
    for k, label in [("salary_ok", "เงินเดือน"), ("level_ok", "ระดับประสบการณ์"), ("work_mode_ok", "รูปแบบงาน")]:
        v = [str(r[k]).strip().upper() for r in data if r.get(k) not in (None, "")]
        ok = sum(x == "Y" for x in v)
        lines.append(f"{label:16} ถูก {ok}/{len(v)} = {ok / len(v):.1%}" if v else f"{label:16} (ยังไม่กรอก)")
    f1 = prf(*allk)[2]
    lines += ["", f"{'✅ ผ่าน' if f1 >= 0.85 else '❌ ไม่ผ่าน'} เกณฑ์ ≥ 85% (F1 ทักษะรวม = {f1:.1%})"]
    if warn:
        lines += ["", "⚠️ ชื่อที่พิมพ์ไม่ตรงกับที่สกัด (ไม่นับ):"] + warn
    out = "\n".join(lines)
    print(out)
    (OUT / "eval_result.txt").write_text(out, encoding="utf-8")
    print(f"\nบันทึก: {OUT / 'eval_result.txt'}")


def main() -> None:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sample")
    s.add_argument("--n", type=int, default=50)
    s.add_argument("--seed", type=float, default=0.42)
    sub.add_parser("score")
    a = ap.parse_args()
    {"sample": cmd_sample, "score": cmd_score}[a.cmd](a)


if __name__ == "__main__":
    main()
