# -*- coding: utf-8 -*-
r"""
จับคู่ทักษะในฐานข้อมูลกับ SkillsFuture Singapore — Unique Skills List

  pip install openpyxl
  python -m extractor.sf_map --file "data\Unique-Skills-List-Sep-2026.xlsx"
  python -m extractor.sf_map --file ... --dry-run       # แสดงผล ไม่เขียน DB

ใช้กฎเข้มงวดเดียวกับ ESCO (ตรงความหมาย ไม่ใช่ตัวสะกดคล้าย) บันทึกลงตาราง skill_standard_map
"""
import argparse
import re
import sys
from collections import defaultdict

import psycopg2

from . import config
from .esco_map import _score, _words
from .normalize import norm_key

STANDARD = "SKILLSFUTURE"
MIN_SCORE = 0.75

# ชื่อของเรา -> ชื่อใน SkillsFuture ที่ตรวจแล้วว่าตรงความหมาย (เติมหลังดูผล audit)
SF_EXACT = {}


NEED = {"skill_title", "skill_id", "skill_type"}


def load_sf(path: str) -> list:
    """เลือกชีตที่หัวตารางมีคอลัมน์ skill_title + skill_id + skill_type และมีแถวมากที่สุด
    (ข้ามชีต Legend ที่มีคำว่า skill_title อยู่ในเนื้อหาคำอธิบาย)"""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    best, best_name = [], None
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        for n, head in enumerate(rows):
            if n > 20:                                   # หัวตารางต้องอยู่ใน 20 แถวแรก
                break
            cols = [str(c).strip().lower() if c is not None else "" for c in head]
            if not NEED.issubset(cols):
                continue
            ix = {c: i for i, c in enumerate(cols) if c}
            out = []
            for r in rows:
                def get(k):
                    i = ix.get(k)
                    return str(r[i]).strip() if i is not None and i < len(r) and r[i] is not None else None
                title = get("skill_title")
                if title:
                    out.append({"id": get("skill_id"), "title": title,
                                "type": get("skill_type"), "desc": get("skill_description")})
            if len(out) > len(best):
                best, best_name = out, ws.title
            break
    if not best:
        print(f"❌ ไม่พบชีตข้อมูล (ชีตที่มี: {[w.title for w in wb.worksheets]})")
        sys.exit(1)
    types = {}
    for it in best:
        types[it["type"]] = types.get(it["type"], 0) + 1
    print(f"📚 SkillsFuture: ชีต '{best_name}' {len(best)} ทักษะ  {types}")
    return best


class SFIndex:
    def __init__(self, items: list):
        self.by_key, self.by_word = {}, defaultdict(list)
        for it in items:
            t = it["title"]
            self.by_key.setdefault(norm_key(t), it)
            base = norm_key(re.sub(r"\s*\(.*?\)", " ", t))
            self.by_key.setdefault(base, it)
            for w in set(_words(t)):
                if len(w) > 2:
                    self.by_word[w].append(it)

    def match(self, name: str):
        if name in SF_EXACT:
            it = self.by_key.get(norm_key(SF_EXACT[name]))
            return (it, 1.0) if it else (None, 0.0)
        k = norm_key(name)
        if k in self.by_key:
            return self.by_key[k], 1.0
        words = [w for w in _words(name) if len(w) > 2]
        if not words:
            return None, 0.0
        # ผู้สมัคร = ทักษะ SF ที่มีคำของเราครบทุกคำ แล้วให้คะแนนด้วยกฎเดียวกับ ESCO
        pool = set(id(x) for x in self.by_word.get(words[0], []))
        for w in words[1:]:
            pool &= set(id(x) for x in self.by_word.get(w, []))
        best, best_s = None, 0.0
        for it in self.by_word.get(words[0], []):
            if id(it) in pool:
                s = _score(name, it["title"])
                if s > best_s:
                    best, best_s = it, s
        return (best, best_s) if best_s >= MIN_SCORE else (None, best_s)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    idx = SFIndex(load_sf(a.file))
    conn = psycopg2.connect(config.DATABASE_URL)
    with conn.cursor() as cur:
        cur.execute("""SELECT s.id, s.name, s.skill_type, count(js.job_id)
                       FROM skills s LEFT JOIN job_skills js ON js.skill_id = s.id
                       GROUP BY s.id ORDER BY count(js.job_id) DESC""")
        skills = cur.fetchall()

    rows, hit_m, tot_m = [], 0, 0
    for sid, name, stype, freq in skills:
        it, sc = idx.match(name)
        tot_m += freq
        if it:
            hit_m += freq
            rows.append((sid, STANDARD, it["id"], it["title"], it["type"], sc, name, freq))

    print(f"\n✅ SkillsFuture จับคู่ได้ {len(rows)}/{len(skills)} ทักษะ "
          f"({len(rows) * 100 // max(len(skills), 1)}%) · ถ่วงตามความถี่ {hit_m * 100 / max(tot_m, 1):.1f}%")
    print("\n▶ ตัวอย่าง 25 อันดับที่พบบ่อย")
    for r in rows[:25]:
        print(f"   {r[6][:28]:30}({r[7]:>4}) -> {r[3][:40]:42} [{r[4] or '-'}]")

    if a.dry_run:
        print("\n--dry-run: ไม่เขียน DB")
        return
    with conn.cursor() as cur:
        cur.execute("DELETE FROM skill_standard_map WHERE standard = %s", (STANDARD,))
        cur.executemany("""INSERT INTO skill_standard_map
                           (skill_id, standard, std_code, std_label, std_type, score)
                           VALUES (%s,%s,%s,%s,%s,%s)""", [r[:6] for r in rows])
        conn.commit()
        cur.execute("SELECT * FROM v_std_coverage_by_standard ORDER BY weighted_pct DESC")
        print(f"\n{'มาตรฐาน':22}{'ทักษะ':>8}{'% ไม่ซ้ำ':>12}{'% ถ่วงน้ำหนัก':>16}")
        for std, n, up, wp in cur.fetchall():
            print(f"{std:22}{n:>8}{up:>11}%{wp:>15}%")
    conn.close()


if __name__ == "__main__":
    main()